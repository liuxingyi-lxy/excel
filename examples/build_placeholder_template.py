"""把 BOOKING.xlsx 的 ='[5]LINK DATA'!D2 这类外部引用替换为 {{D2}} 形式，
生成一份「写满占位符」的样本模板，方便不看外部引用时也能用同一份映射 JSON 走通渲染。
"""

import re
import shutil
import sys
from pathlib import Path

import openpyxl

SRC = Path("/home/gem/.aily/workdir/task_7698581490505878806/artifacts/BOOKING.xlsx")
DST = Path("/home/gem/.aily/workdir/task_7698581490505878806/artifacts/excel-template-exporter/examples/BOOKING-template-placeholder.xlsx")

REF = re.compile(r"='\[[^\]]*\][^']+'\s*!\s*(\$?[A-Z]{1,3}\$?\d+)")


def normalize_ref(ref: str) -> str:
    return ref.replace("$", "").upper()


def main():
    if not SRC.exists():
        print(f"找不到 {SRC}", file=sys.stderr)
        sys.exit(1)
    DST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SRC, DST)

    wb = openpyxl.load_workbook(DST)
    replaced = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and REF.search(v):
                    new = REF.sub(lambda m: "{{" + normalize_ref(m.group(1)) + "}}", v)
                    if new != v:
                        c.value = new
                        replaced += 1
    wb.save(DST)
    print(f"{SRC.name} -> {DST.name}  替换 {replaced} 个公式格子为 {{占位符}}")


if __name__ == "__main__":
    main()
